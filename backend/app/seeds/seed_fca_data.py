import os
from datetime import datetime
import openpyxl
from sqlalchemy.orm import Session
from app.core.database import SessionLocal, Base, engine
from app.core.security import hash_password
from app.models.models import User, Role, FCAgent, FCAMonthlyRecord

def find_file(filename: str):
    candidates = [
        filename,
        os.path.join("..", filename),
        os.path.join(os.path.dirname(__file__), "..", "..", "..", filename),
        os.path.join(os.getcwd(), filename),
        os.path.join(os.getcwd(), "..", filename),
    ]
    for c in candidates:
        if os.path.exists(c):
            return os.path.abspath(c)
    return None

MONTH_LABELS = {
    "2026-01": "Jan 2026",
    "2026-02": "Feb 2026",
    "2026-03": "Mar 2026",
    "2026-04": "Apr 2026",
    "2026-05": "May 2026",
    "2026-06": "Jun 2026",
    "2026-07": "Jul 2026",
    "2026-08": "Aug 2026",
    "2026-09": "Sep 2026",
    "2026-10": "Oct 2026",
    "2026-11": "Nov 2026",
    "2026-12": "Dec 2026",
}

def seed_fca_data(db: Session):
    print("--- Seeding FCA and BVS Tracking Data ---")
    Base.metadata.create_all(bind=engine)

    # 1. Ensure Shakeel Ahmad user exists
    staff_role = db.query(Role).filter(Role.name == "Staff").first()
    if not staff_role:
        staff_role = Role(name="Staff", description="Operations & Staff Member")
        db.add(staff_role)
        db.flush()

    shakeel = db.query(User).filter(User.email == "shakeel@pos.com").first()
    if not shakeel:
        shakeel = User(
            name="Shakeel Ahmad",
            email="shakeel@pos.com",
            hashed_password=hash_password("posUfone@123"),
            role_id=staff_role.id,
            phone="+92 333 9876543",
            is_active=True
        )
        db.add(shakeel)
        db.flush()
        print("Created user Shakeel Ahmad (shakeel@pos.com)")
    else:
        shakeel.hashed_password = hash_password("posUfone@123")
        shakeel.is_active = True

    # Check if already seeded with records
    existing_count = db.query(FCAgent).count()
    if existing_count > 0:
        print(f"FCAgent already contains {existing_count} records. Verifying...")
        db.commit()
        return

    f1 = find_file("FCA Table AUG 2026-1.xlsx")
    if not f1:
        print("Warning: FCA Table AUG 2026-1.xlsx not found.")
        db.commit()
        return

    wb1 = openpyxl.load_workbook(f1, data_only=True)
    agents_data = {}  # bvs_id -> dict

    # --- SHEET 1 ---
    ws1 = wb1['Sheet1']
    cat1 = 'SABIR RSO MARKET'
    for r in range(3, ws1.max_row + 1):
        name = ws1.cell(r, 2).value
        market = ws1.cell(r, 3).value
        bvs = ws1.cell(r, 4).value
        if name and 'SABIR RSO' in str(name):
            cat1 = 'RIAZ RSO MARKET'
            continue
        if name and 'RIAZ RSO' in str(name):
            cat1 = 'KHIZAR ALI RSO MARKET'
            continue
        if name and 'KHIZAR' in str(name) and not bvs:
            continue
        if bvs:
            bvs_clean = str(bvs).strip().upper()
            months = {}
            for c, m in enumerate(['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06', '2026-07', '2026-08'], 5):
                val = ws1.cell(r, c).value
                months[m] = int(val) if val is not None and isinstance(val, (int, float)) else 0
            agents_data[bvs_clean] = {
                'bvs_id': bvs_clean,
                'name': str(name).strip() if name else '',
                'market': str(market).strip() if market else '',
                'category': cat1,
                'channel': 'Market FCA',
                'months': months
            }

    # --- SHEET 2 ---
    ws2 = wb1['Sheet2']
    cat2 = 'MAAZ RSO MARKET'
    for r in range(4, ws2.max_row + 1):
        name = ws2.cell(r, 3).value
        market = ws2.cell(r, 4).value
        bvs = ws2.cell(r, 5).value
        if name and 'MAAZ' in str(name):
            cat2 = 'FRANCHISE OFFICE'
            continue
        if name and 'FRANCHISE' in str(name) and not bvs:
            cat2 = 'DSO s'
            continue
        if name and 'DSO s' in str(name):
            cat2 = 'BDO'
            continue
        if (market and 'TKB' in str(market)) or (name and 'TKB' in str(name)):
            cat2 = 'TKB FCA'
            if not bvs:
                continue
        if bvs:
            bvs_clean = str(bvs).strip().upper()
            months = {}
            for c, m in enumerate(['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06', '2026-07', '2026-08'], 6):
                val = ws2.cell(r, c).value
                months[m] = int(val) if val is not None and isinstance(val, (int, float)) else 0
            
            channel = 'DSO' if 'DSO' in cat2 else ('Office' if 'OFFICE' in cat2 or 'FRANCHISE' in cat2 else ('BDO' if 'BDO' in cat2 else 'Market FCA'))
            if bvs_clean in agents_data:
                for mk, mv in months.items():
                    if mv > 0 or mk not in agents_data[bvs_clean]['months']:
                        agents_data[bvs_clean]['months'][mk] = mv
            else:
                agents_data[bvs_clean] = {
                    'bvs_id': bvs_clean,
                    'name': str(name).strip() if name else (str(market).strip() if market else 'BVS Agent'),
                    'market': str(market).strip() if market else '',
                    'category': cat2,
                    'channel': channel,
                    'months': months
                }

    # --- SHEET 3 ---
    ws3 = wb1['Sheet3']
    for r in range(3, 12):
        name = ws3.cell(r, 3).value
        market = ws3.cell(r, 4).value
        bvs = ws3.cell(r, 5).value
        val = ws3.cell(r, 6).value
        if bvs:
            bvs_clean = str(bvs).strip().upper()
            if bvs_clean in agents_data:
                if val is not None and int(val) > agents_data[bvs_clean]['months'].get('2026-08', 0):
                    agents_data[bvs_clean]['months']['2026-08'] = int(val)
            else:
                agents_data[bvs_clean] = {
                    'bvs_id': bvs_clean,
                    'name': str(name).strip() if name else 'Franchise Agent',
                    'market': str(market).strip() if market else 'Dargai',
                    'category': 'FRANCHISE BVS',
                    'channel': 'Office',
                    'months': {'2026-08': int(val) if val is not None else 0}
                }

    # --- PROCESS BVS SEP 2026 FCA ---
    f2 = find_file("BVS  SEP 2026 FCA.xlsx")
    if f2:
        wb2 = openpyxl.load_workbook(f2, data_only=True)
        ws_sep = wb2['Sheet1']
        for r in range(2, ws_sep.max_row + 1):
            bvs = ws_sep.cell(r, 1).value
            val = ws_sep.cell(r, 2).value
            if bvs:
                bvs_clean = str(bvs).strip().upper()
                count = int(val) if val is not None and isinstance(val, (int, float)) else 0
                target_key = bvs_clean
                if target_key not in agents_data:
                    for k in agents_data.keys():
                        if k.startswith(target_key) or target_key.startswith(k):
                            target_key = k
                            break
                if target_key in agents_data:
                    agents_data[target_key]['months']['2026-09'] = count
                else:
                    agents_data[bvs_clean] = {
                        'bvs_id': bvs_clean,
                        'name': f'Agent {bvs_clean}',
                        'market': 'Dargai Sector',
                        'category': 'MARKET FCA',
                        'channel': 'Market FCA',
                        'months': {'2026-09': count}
                    }

    # Insert into database
    for bvs_id, data in agents_data.items():
        agent = FCAgent(
            bvs_id=data['bvs_id'],
            name=data['name'] or f"Agent {bvs_id}",
            market=data['market'] or "Dargai",
            category=data['category'] or "Market FCA",
            channel=data['channel'] or "Market FCA",
            status="Active"
        )
        db.add(agent)
        db.flush()

        for m_key, count in data['months'].items():
            rec = FCAMonthlyRecord(
                agent_id=agent.id,
                month_key=m_key,
                month_label=MONTH_LABELS.get(m_key, m_key),
                sims_sold=count,
                source="Initial Import (Excel)",
                updated_by_user_id=shakeel.id
            )
            db.add(rec)

    db.commit()
    print(f"Successfully seeded {len(agents_data)} FCA Agents with complete monthly records!")

if __name__ == "__main__":
    db = SessionLocal()
    seed_fca_data(db)
    db.close()
