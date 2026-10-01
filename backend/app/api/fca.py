import io
from datetime import datetime, date
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.database import get_db
from app.api.auth import get_current_user
from app.models.models import FCAgent, FCAMonthlyRecord, User, AuditLog

router = APIRouter(prefix="/fca", tags=["FCA & BVS Performance Tracking"])

MONTH_LABELS = {
    "2026-01": "Jan 2026",
    "2026-02": "Feb 2026",
    "2026-03": "Mar 2026",
    "2026-04": "Apr 2026",
    "2026-05": "May 2026",
    "2026-06": "Jun 2026",
    "2026-07": "Jul 2026",
    "2026-08": "Aug 2026",
    "2026-09": "Sep 2026",
    "2026-10": "Oct 2026",
    "2026-11": "Nov 2026",
    "2026-12": "Dec 2026",
    "2027-01": "Jan 2027",
    "2027-02": "Feb 2027",
    "2027-03": "Mar 2027",
}

def get_label(m_key: str) -> str:
    if m_key in MONTH_LABELS:
        return MONTH_LABELS[m_key]
    try:
        parts = m_key.split("-")
        dt = datetime(int(parts[0]), int(parts[1]), 1)
        return dt.strftime("%b %Y")
    except Exception:
        return m_key

@router.get("/agents")
def list_fca_agents(
    search: Optional[str] = None,
    category: Optional[str] = None,
    channel: Optional[str] = None,
    sort_by: Optional[str] = "category",
    sort_dir: Optional[str] = "asc",
    db: Session = Depends(get_db)
):
    query = db.query(FCAgent)

    if category and category != "All":
        query = query.filter(FCAgent.category == category)
    if channel and channel != "All":
        query = query.filter(FCAgent.channel == channel)
    if search:
        s = f"%{search.strip()}%"
        query = query.filter(
            (FCAgent.bvs_id.ilike(s)) |
            (FCAgent.name.ilike(s)) |
            (FCAgent.market.ilike(s)) |
            (FCAgent.category.ilike(s))
        )

    agents = query.all()

    # Get all distinct month keys in the database
    month_records = db.query(FCAMonthlyRecord.month_key).distinct().order_by(FCAMonthlyRecord.month_key.asc()).all()
    all_month_keys = [m[0] for m in month_records if m[0]]
    if not all_month_keys:
        all_month_keys = ["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09"]
    else:
        all_month_keys = sorted(list(set(all_month_keys)))

    all_months = [{"key": mk, "label": get_label(mk)} for mk in all_month_keys]

    # Distinct categories for filtering
    cat_records = db.query(FCAgent.category).distinct().all()
    categories = sorted([c[0] for c in cat_records if c[0]])

    results = []
    total_sims_all_time = 0
    latest_month_key = all_month_keys[-1] if all_month_keys else "2026-09"
    latest_month_total = 0

    for a in agents:
        m_dict = {r.month_key: r.sims_sold for r in a.monthly_records}
        agent_total = sum(m_dict.values())
        active_months = len([v for v in m_dict.values() if v > 0])
        avg = round(agent_total / max(1, len(all_month_keys)), 1)
        total_sims_all_time += agent_total
        latest_month_total += m_dict.get(latest_month_key, 0)

        results.append({
            "id": a.id,
            "bvs_id": a.bvs_id,
            "name": a.name or f"Agent {a.bvs_id}",
            "market": a.market or "Dargai",
            "category": a.category or "Market FCA",
            "channel": a.channel or "Market FCA",
            "status": a.status or "Active",
            "months": m_dict,
            "total_sims": agent_total,
            "monthly_avg": avg,
            "active_months_count": active_months,
            "latest_month_sales": m_dict.get(latest_month_key, 0)
        })

    # Sorting
    reverse = (sort_dir.lower() == "desc")
    if sort_by == "total":
        results.sort(key=lambda x: x["total_sims"], reverse=reverse)
    elif sort_by == "latest":
        results.sort(key=lambda x: x["latest_month_sales"], reverse=reverse)
    elif sort_by == "bvs_id":
        results.sort(key=lambda x: x["bvs_id"], reverse=reverse)
    elif sort_by == "name":
        results.sort(key=lambda x: x["name"], reverse=reverse)
    elif sort_by in all_month_keys:
        results.sort(key=lambda x: x["months"].get(sort_by, 0), reverse=reverse)
    else:
        results.sort(key=lambda x: (x["category"], x["name"]), reverse=reverse)

    return {
        "agents": results,
        "all_months": all_months,
        "categories": categories,
        "summary": {
            "total_agents": len(results),
            "total_sims_all_time": total_sims_all_time,
            "active_bvs_count": len([r for r in results if r["status"] == "Active"]),
            "latest_month_key": latest_month_key,
            "latest_month_label": get_label(latest_month_key),
            "latest_month_total": latest_month_total
        }
    }

@router.get("/summary")
def get_fca_summary(db: Session = Depends(get_db)):
    agents = db.query(FCAgent).all()
    records = db.query(FCAMonthlyRecord).all()

    # Monthly totals
    month_totals: Dict[str, int] = {}
    for r in records:
        month_totals[r.month_key] = month_totals.get(r.month_key, 0) + r.sims_sold

    sorted_months = sorted(month_totals.keys())
    monthly_trend = [
        {
            "month_key": mk,
            "month_label": get_label(mk),
            "total_sims": month_totals[mk]
        }
        for mk in sorted_months
    ]

    # Category breakdown
    cat_totals: Dict[str, Dict[str, Any]] = {}
    for a in agents:
        c = a.category or "Other"
        if c not in cat_totals:
            cat_totals[c] = {"category": c, "agent_count": 0, "total_sims": 0}
        cat_totals[c]["agent_count"] += 1
        cat_totals[c]["total_sims"] += sum(r.sims_sold for r in a.monthly_records)

    return {
        "total_registered_bvs": len(agents),
        "total_active_bvs": len([a for a in agents if a.status == "Active"]),
        "monthly_trend": monthly_trend,
        "categories": list(cat_totals.values()),
        "latest_month": monthly_trend[-1] if monthly_trend else None
    }

@router.post("/agents")
def create_fca_agent(
    payload: Dict[str, Any],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    bvs_id = payload.get("bvs_id", "").strip().upper()
    if not bvs_id:
        raise HTTPException(status_code=400, detail="BVS ID is required")

    existing = db.query(FCAgent).filter(FCAgent.bvs_id == bvs_id).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"BVS ID {bvs_id} already exists")

    agent = FCAgent(
        bvs_id=bvs_id,
        name=payload.get("name", "").strip() or f"Agent {bvs_id}",
        market=payload.get("market", "").strip() or "Dargai",
        category=payload.get("category", "Market FCA").strip(),
        channel=payload.get("channel", "Market FCA").strip(),
        status=payload.get("status", "Active")
    )
    db.add(agent)
    db.flush()

    # Optional initial months
    months_data = payload.get("months", {})
    for mk, val in months_data.items():
        if val is not None:
            rec = FCAMonthlyRecord(
                agent_id=agent.id,
                month_key=mk,
                month_label=get_label(mk),
                sims_sold=int(val),
                source="Manual Entry",
                updated_by_user_id=current_user.id
            )
            db.add(rec)

    # Audit log
    audit = AuditLog(
        user_id=current_user.id,
        user_name=current_user.name,
        action="Create",
        entity="FCAgent",
        entity_id=bvs_id,
        new_value=f"Created BVS Agent: {agent.name} ({bvs_id}) in {agent.market}"
    )
    db.add(audit)
    db.commit()

    return {"message": "Agent created successfully", "agent_id": agent.id, "bvs_id": bvs_id}

@router.put("/agents/{agent_id}")
def update_fca_agent(
    agent_id: int,
    payload: Dict[str, Any],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    agent = db.query(FCAgent).filter(FCAgent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agent not found")

    if "name" in payload:
        agent.name = payload["name"].strip()
    if "market" in payload:
        agent.market = payload["market"].strip()
    if "category" in payload:
        agent.category = payload["category"].strip()
    if "channel" in payload:
        agent.channel = payload["channel"].strip()
    if "status" in payload:
        agent.status = payload["status"].strip()

    # Update month values if provided
    if "months" in payload and isinstance(payload["months"], dict):
        for mk, val in payload["months"].items():
            if val is not None:
                rec = db.query(FCAMonthlyRecord).filter(
                    FCAMonthlyRecord.agent_id == agent.id,
                    FCAMonthlyRecord.month_key == mk
                ).first()
                if rec:
                    rec.sims_sold = int(val)
                    rec.updated_by_user_id = current_user.id
                else:
                    rec = FCAMonthlyRecord(
                        agent_id=agent.id,
                        month_key=mk,
                        month_label=get_label(mk),
                        sims_sold=int(val),
                        source="Manual Update",
                        updated_by_user_id=current_user.id
                    )
                    db.add(rec)

    db.commit()
    return {"message": "Agent updated successfully"}

@router.post("/agents/{agent_id}/months")
def update_agent_month_sale(
    agent_id: int,
    payload: Dict[str, Any],
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    agent = db.query(FCAgent).filter(FCAgent.id == agent_id).first()
    if not agent:
        raise HTTPException(status_code=404, detail="Agent not found")

    month_key = payload.get("month_key")
    if not month_key:
        raise HTTPException(status_code=400, detail="month_key is required")

    sims_sold = int(payload.get("sims_sold", 0))

    rec = db.query(FCAMonthlyRecord).filter(
        FCAMonthlyRecord.agent_id == agent.id,
        FCAMonthlyRecord.month_key == month_key
    ).first()

    old_val = 0
    if rec:
        old_val = rec.sims_sold
        rec.sims_sold = sims_sold
        rec.updated_by_user_id = current_user.id
        rec.updated_at = datetime.utcnow()
    else:
        rec = FCAMonthlyRecord(
            agent_id=agent.id,
            month_key=month_key,
            month_label=get_label(month_key),
            sims_sold=sims_sold,
            source="Manual Entry",
            updated_by_user_id=current_user.id
        )
        db.add(rec)

    audit = AuditLog(
        user_id=current_user.id,
        user_name=current_user.name,
        action="Update",
        entity="FCAMonthlyRecord",
        entity_id=f"{agent.bvs_id}:{month_key}",
        old_value=str(old_val),
        new_value=str(sims_sold)
    )
    db.add(audit)
    db.commit()

    return {"message": "Monthly record updated", "month_key": month_key, "sims_sold": sims_sold}

@router.post("/upload-monthly-excel")
async def upload_monthly_excel(
    file: UploadFile = File(...),
    month_key: Optional[str] = Form(None),
    month_label: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Parses an uploaded Excel file (matching BVS SEP 2026 FCA format),
    matches BVS IDs, and updates or creates monthly progress entries.
    """
    if not file.filename.endswith(('.xlsx', '.xlsm', '.xls')):
        raise HTTPException(status_code=400, detail="Only Excel (.xlsx) files are supported")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    ws = wb.active

    # 1. Detect target month if not provided
    detected_month_key = month_key
    detected_month_label = month_label

    if not detected_month_key:
        # Check cell (1, 2) or header
        header_val = ws.cell(1, 2).value
        if isinstance(header_val, (datetime, date)):
            detected_month_key = header_val.strftime("%Y-%m")
            detected_month_label = header_val.strftime("%b %Y")
        elif isinstance(header_val, str) and len(header_val) > 2:
            detected_month_label = header_val.strip()
            detected_month_key = "2026-09"  # fallback default
        else:
            detected_month_key = "2026-09"
            detected_month_label = "Sep 2026"

    if not detected_month_label:
        detected_month_label = get_label(detected_month_key)

    # 2. Build map of existing agents: bvs_id -> agent
    existing_agents = db.query(FCAgent).all()
    agent_map: Dict[str, FCAgent] = {}
    for a in existing_agents:
        agent_map[a.bvs_id.strip().upper()] = a

    matched_count = 0
    new_agents_count = 0
    total_sims = 0
    updated_records = []

    # 3. Iterate rows (from row 2 downwards)
    for r in range(2, ws.max_row + 1):
        bvs_raw = ws.cell(r, 1).value
        val_raw = ws.cell(r, 2).value
        if not bvs_raw:
            continue

        bvs_clean = str(bvs_raw).strip().upper()
        count = int(val_raw) if val_raw is not None and isinstance(val_raw, (int, float)) else 0
        total_sims += count

        # Match exact or normalized prefix
        target_agent = agent_map.get(bvs_clean)
        if not target_agent:
            # check prefix / fuzzy match
            for k, ag in agent_map.items():
                if k.startswith(bvs_clean) or bvs_clean.startswith(k):
                    target_agent = ag
                    break

        if target_agent:
            # Update or create monthly record
            rec = db.query(FCAMonthlyRecord).filter(
                FCAMonthlyRecord.agent_id == target_agent.id,
                FCAMonthlyRecord.month_key == detected_month_key
            ).first()
            if rec:
                rec.sims_sold = count
                rec.source = f"Excel Upload ({file.filename})"
                rec.updated_by_user_id = current_user.id
                rec.updated_at = datetime.utcnow()
            else:
                rec = FCAMonthlyRecord(
                    agent_id=target_agent.id,
                    month_key=detected_month_key,
                    month_label=detected_month_label,
                    sims_sold=count,
                    source=f"Excel Upload ({file.filename})",
                    updated_by_user_id=current_user.id
                )
                db.add(rec)
            matched_count += 1
            updated_records.append({
                "bvs_id": target_agent.bvs_id,
                "name": target_agent.name,
                "sims": count,
                "status": "Updated"
            })
        else:
            # Auto-create new agent
            new_agent = FCAgent(
                bvs_id=bvs_clean,
                name=f"Agent {bvs_clean}",
                market="Dargai Sector",
                category="MARKET FCA",
                channel="Market FCA",
                status="Active"
            )
            db.add(new_agent)
            db.flush()
            agent_map[bvs_clean] = new_agent

            rec = FCAMonthlyRecord(
                agent_id=new_agent.id,
                month_key=detected_month_key,
                month_label=detected_month_label,
                sims_sold=count,
                source=f"Excel Upload ({file.filename})",
                updated_by_user_id=current_user.id
            )
            db.add(rec)
            new_agents_count += 1
            updated_records.append({
                "bvs_id": bvs_clean,
                "name": new_agent.name,
                "sims": count,
                "status": "New Agent Added"
            })

    # Record Audit Log
    audit = AuditLog(
        user_id=current_user.id,
        user_name=current_user.name,
        action="Excel Upload",
        entity="FCAMonthlyProgress",
        entity_id=detected_month_key,
        new_value=f"Uploaded monthly SIMs for {detected_month_label}: {matched_count} updated, {new_agents_count} new agents. Total SIMs: {total_sims}"
    )
    db.add(audit)
    db.commit()

    return {
        "success": True,
        "message": f"Successfully updated {matched_count} BVS agents for {detected_month_label}! ({new_agents_count} new agents added).",
        "month_key": detected_month_key,
        "month_label": detected_month_label,
        "matched_count": matched_count,
        "new_agents_count": new_agents_count,
        "total_sims_recorded": total_sims,
        "filename": file.filename
    }

@router.get("/export")
def export_fca_excel(db: Session = Depends(get_db)):
    """Generates an executive Excel workbook with all agents and all monthly columns."""
    agents = db.query(FCAgent).order_by(FCAgent.category.asc(), FCAgent.name.asc()).all()
    month_records = db.query(FCAMonthlyRecord.month_key).distinct().order_by(FCAMonthlyRecord.month_key.asc()).all()
    all_month_keys = sorted(list(set([m[0] for m in month_records if m[0]])))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FCA Monthly Progress"

    # Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    sub_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Title
    ws.merge_cells("A1:K1")
    ws["A1"] = "UFONE FRANCHISE DARGAI - FCA & BVS MONTHLY SIMS SALES PROGRESS"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Master Performance Ledger across all Market FCAs, DSOs, and Office BVS IDs | Generated {datetime.utcnow().strftime('%d-%b-%Y')}"
    ws["A2"].font = sub_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Table Headers
    base_headers = ["#", "BVS ID", "Agent Name", "Market / Route", "Category", "Channel"]
    month_headers = [get_label(mk) for mk in all_month_keys]
    final_headers = base_headers + month_headers + ["Total SIMs", "Monthly Avg"]

    header_row = 4
    for c_idx, h_text in enumerate(final_headers, 1):
        cell = ws.cell(row=header_row, column=c_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if c_idx > 4 else "left", vertical="center")
        cell.border = thin_border

    # Rows
    current_row = 5
    for idx, a in enumerate(agents, 1):
        m_dict = {r.month_key: r.sims_sold for r in a.monthly_records}
        row_vals = [
            idx,
            a.bvs_id,
            a.name,
            a.market,
            a.category,
            a.channel
        ]
        agent_total = 0
        for mk in all_month_keys:
            val = m_dict.get(mk, 0)
            row_vals.append(val)
            agent_total += val

        avg = round(agent_total / max(1, len(all_month_keys)), 1)
        row_vals.append(agent_total)
        row_vals.append(avg)

        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = bold_font if c_idx in [2, len(row_vals) - 1] else regular_font
            cell.border = thin_border
            if c_idx > 6:
                cell.alignment = Alignment(horizontal="right")
        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"FCA_Monthly_Progress_Dargai_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
